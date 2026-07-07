import json, re
from pathlib import Path

import fitz
import pytesseract
from PIL import Image, ImageEnhance, ImageFilter
from docx import Document as DocxDocument
import openpyxl
from google import genai
from google.genai import types
import streamlit as st

from config import GEMINI_API_KEY, MODEL_NAME, LITE_MODEL_NAME, IMAGE_EXT

def preprocess(img):
    if img.mode != "RGB": img = img.convert("RGB")
    w, h = img.size
    if w < 1000: img = img.resize((1000, int(h * 1000 / w)), Image.LANCZOS)
    img = img.filter(ImageFilter.SHARPEN)
    img = ImageEnhance.Contrast(img).enhance(1.5)
    return img.convert("L")


def extract_text(path):
    ext = Path(path).suffix.lower()
    if ext in IMAGE_EXT:
        img = preprocess(Image.open(path))
        return max(
            [pytesseract.image_to_string(img, config=c) for c in ["--psm 3", "--psm 6", "--psm 11"]],
            key=len
        )
    if ext == ".pdf":
        doc = fitz.open(path)
        pages = []
        for p in doc:
            t = p.get_text("text").strip()
            if len(t) > 50:
                pages.append(t)
            else:
                pix = p.get_pixmap(dpi=250)
                img = preprocess(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))
                pages.append(pytesseract.image_to_string(img))
        doc.close()
        return "\n--- PAGE BREAK ---\n".join(pages)
    if ext == ".docx":
        return "\n".join(p.text for p in DocxDocument(path).paragraphs if p.text.strip())
    if ext == ".txt":
        with open(path) as f:
            return f.read()
    if ext in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(path, data_only=True)
        return "\n".join(
            "  |  ".join(str(c) for c in r if c)
            for s in wb.worksheets
            for r in s.iter_rows(values_only=True)
        )
    raise ValueError(f"Unsupported: {ext}")


def _llm_vision_call(system_prompt, path, max_tokens=8000):
    """Vision call - renders PDF pages as images and sends to Gemini Flash.
    Used for resumes and sea service docs where table layout matters."""
    doc = fitz.open(path)
    image_parts = []
    for page in doc:
        pix = page.get_pixmap(dpi=150)  # 150 Dpi = 6 tiles/page = ~1548 tokens/page
        image_parts.append(
            types.Part.from_bytes(data=pix.tobytes("png"), mime_type="image/png")
        )
    doc.close()

    if not image_parts:
        return None

    client = genai.Client(api_key=GEMINI_API_KEY)
    for attempt in range(3):
        try:
            response = client.models.generate_content(
                model=MODEL_NAME,
                contents=[
                    types.Content(
                        role="user",
                        parts=image_parts + [types.Part(text="Extract all information from these document pages.")]
                    )
                ],
                config=types.GenerateContentConfig(
                    system_instruction=system_prompt,
                    temperature=0.1,
                    max_output_tokens=max_tokens,
                    thinking_config=types.ThinkingConfig(thinking_budget=0),
                    response_mime_type="application/json",
                )
            )
            raw = response.text.strip()
            if not raw: continue
            raw = re.sub(r"^```[a-z]*\n?", "", raw)
            raw = re.sub(r"\n?```$", "", raw)
            start = raw.find("{"); end = raw.rfind("}") + 1
            if start == -1 or end == 0: continue
            return json.loads(raw[start:end])
        except json.JSONDecodeError:
            continue
        except Exception as e:
            st.error(f"API error: {e}")
            return None
    st.error("Pipeline failed after 3 attempts.")
    return None


def _llm_call(system_prompt, user_text, max_tokens=4000, text_limit=30000, model=None):
    """Core LLM caller - shared by all pipelines."""
    model = model or MODEL_NAME
    if len(user_text) > text_limit:
        user_text = user_text[:text_limit]
    if not user_text.strip():
        return None

    client = genai.Client(api_key=GEMINI_API_KEY)

    for attempt in range(3):
        try:
            response = client.models.generate_content(
                model=model,
                contents=f"Extract all information:\n\n{user_text}",
                config=types.GenerateContentConfig(
                    system_instruction=system_prompt,
                    temperature=0.1,
                    max_output_tokens=max_tokens,
                    thinking_config=types.ThinkingConfig(thinking_budget=0),
                    response_mime_type="application/json",
                )
            )
            raw = response.text.strip()
            if not raw: continue
            raw = re.sub(r"^```[a-z]*\n?", "", raw)
            raw = re.sub(r"\n?```$", "", raw)
            start = raw.find("{"); end = raw.rfind("}") + 1
            if start == -1 or end == 0: continue
            return json.loads(raw[start:end])
        except json.JSONDecodeError:
            continue
        except Exception as e:
            st.error(f"API error: {e}")
            return None
    st.error("Pipeline failed after 3 attempts.")
    return None


def classify_document(text):
    """Step 1 for Master Upload - lightweight classifier, ~200 tokens."""
    return _llm_call(CLASSIFIER_PROMPT, text, max_tokens=200, text_limit=3000, model=LITE_MODEL_NAME)


# Sections needing Flash on the text path (complex tables)
_FLASH_SECTIONS = {"Sea Service Experience"}

# Sections that use Gemini vision even on digital PDFs (layout/table critical)
_VISION_SECTIONS = {"Sea Service Experience"}


def call_llm(text, target_section=None, is_vessel=False, is_resume=False, path=None):
    """
    Unified entry point.
    - target_section set  -> use section-specific prompt (1 call)
    - is_vessel = True    -> use vessel prompt (1 call)
    - is_resume = True    -> use resume prompt (1 call)
    - path                -> if PDF, enables vision routing for eligible doc types
    - all None/False      -> full resume prompt as fallback

    Vision routing (path must be a .pdf):
    - Resumes always use vision - sea service tables need visual layout
    - Sea service section upload also uses vision
    - All other sections use text path (flat structure, PyMuPDF/Tesseract sufficient)

    Model routing (text path):
    - Flash (MODEL_NAME): resumes, sea service experience
    - Flash-Lite (LITE_MODEL_NAME): classify, vessel certs, all flat sections
    """
    use_vision = path and Path(path).suffix.lower() == ".pdf"

    if is_vessel:
        return _llm_call(VESSEL_PROMPT, text, max_tokens=2000, model=LITE_MODEL_NAME)
    if is_resume:
        if use_vision:
            return _llm_vision_call(RESUME_PROMPT, path, max_tokens=8000)
        return _llm_call(RESUME_PROMPT, text, max_tokens=8000, text_limit=50000, model=MODEL_NAME)
    if target_section and target_section in SECTION_PROMPTS:
        if use_vision and target_section in _VISION_SECTIONS:
            return _llm_vision_call(SECTION_PROMPTS[target_section], path, max_tokens=4000)
        m = MODEL_NAME if target_section in _FLASH_SECTIONS else LITE_MODEL_NAME
        return _llm_call(SECTION_PROMPTS[target_section], text, max_tokens=4000, model=m)
    return _llm_call(RESUME_PROMPT, text, max_tokens=8000, model=MODEL_NAME)
